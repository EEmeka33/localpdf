"""Convert from PDF: images, Word, PowerPoint, Excel, PDF/A, text, HTML, Markdown."""
from __future__ import annotations

import io
import shutil
from pathlib import Path

import pymupdf as fitz

from ..registry import Opt, tool
from ..tables import page_layout
from ..utils import (
    ToolError, ensure_dir, ghostscript, open_pdf, parse_pages, reading_order, run, soffice_convert, tempdir, which,
)


@tool(
    name="pdf-to-jpg", title="PDF to JPG", category="Convert from PDF",
    help="Convert each PDF page into a JPG (or PNG/WEBP/TIFF) or extract all images contained in a PDF.",
    aliases=("pdf-to-image", "pdf-to-png"), output="dir",
    opts=[
        Opt("mode", "choice", "pages", "Render pages, or extract the embedded images.", ["pages", "extract"]),
        Opt("format", "choice", "jpg", "Image format for rendered pages.", ["jpg", "png", "webp", "tiff"]),
        Opt("dpi", "int", 150, "Resolution (72 = screen, 150 = normal, 300 = print)."),
        Opt("pages", "str", "all", "Pages to convert, e.g. '1-3,7'."),
        Opt("quality", "int", 90, "JPEG / WEBP quality."),
        Opt("min_size", "int", 50, "Extract mode: skip images smaller than this many pixels (icons, lines)."),
    ],
)
def pdf_to_jpg(inputs, output: Path, mode="pages", format="jpg", dpi=150, pages="all", quality=90, min_size=50):
    doc = open_pdf(inputs[0])
    stem = Path(inputs[0]).stem
    ensure_dir(output)
    written = []
    if mode == "extract":
        seen = set()
        for pno in parse_pages(pages, doc.page_count):
            for img in doc[pno].get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                info = doc.extract_image(xref)
                if not info or min(info["width"], info["height"]) < min_size:
                    continue
                data, ext = info["image"], info["ext"]
                if info.get("smask"):  # re-attach transparency
                    pix = fitz.Pixmap(fitz.Pixmap(doc, xref), fitz.Pixmap(doc, info["smask"]))
                    data, ext = pix.tobytes("png"), "png"
                target = output / f"{stem}_p{pno + 1}_img{xref}.{ext}"
                target.write_bytes(data)
                written.append(target)
        if not written:
            raise ToolError("no images found in this PDF")
        return written

    from PIL import Image

    width = len(str(doc.page_count))
    for pno in parse_pages(pages, doc.page_count):
        pix = doc[pno].get_pixmap(dpi=dpi, alpha=False)
        target = output / f"{stem}_{str(pno + 1).zfill(width)}.{format}"
        if format == "png":
            pix.save(target)
        else:
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            fmt = {"jpg": "JPEG", "webp": "WEBP", "tiff": "TIFF"}[format]
            img.save(target, fmt, quality=quality, dpi=(dpi, dpi), **({"compression": "tiff_lzw"} if fmt == "TIFF" else {}))
        written.append(target)
    return written


@tool(
    name="pdf-to-word", title="PDF to Word", category="Convert from PDF",
    help="Easily convert your PDF files into easy to edit DOCX documents.",
    aliases=("pdf-to-docx",), output="docx", requires=("pdf2docx", "soffice"),
    opts=[
        Opt("engine", "choice", "auto", "Converter: pdf2docx (best layout) or LibreOffice.", ["auto", "pdf2docx", "libreoffice"]),
        Opt("pages", "str", "all", "Pages to convert."),
        Opt("format", "choice", "docx", "Output format.", ["docx", "odt", "rtf", "doc"]),
        Opt("tables", "choice", "auto", "Borderless-table detection. It can take minutes per page on dense multi-column "
            "text, so 'auto' gives it 8 s per page and otherwise converts without it.", ["auto", "on", "off"]),
    ],
)
def pdf_to_word(inputs, output: Path, engine="auto", pages="all", format="docx", tables="auto"):
    src = Path(inputs[0])
    with tempdir() as tmp:
        if pages != "all":
            doc = open_pdf(src)
            doc.select(parse_pages(pages, doc.page_count))
            doc.save(tmp / src.name)
            src = tmp / src.name
        docx_path = tmp / f"{src.stem}.docx"
        done = False
        if engine in ("auto", "pdf2docx"):
            try:
                import pdf2docx  # noqa: F401
            except ImportError:
                if engine == "pdf2docx":
                    raise ToolError("pip install pdf2docx") from None
            else:
                if tables == "auto":
                    budget = 8 * fitz.open(src).page_count + 5
                    if not _pdf2docx_with_timeout(src, docx_path, budget):
                        print(f"(borderless-table detection exceeded {budget}s – converting without it)")
                        _pdf2docx(src, docx_path, False)
                else:
                    _pdf2docx(src, docx_path, tables == "on")
                done = True
        if not done:
            docx_path = soffice_convert(src, tmp, "docx:MS Word 2007 XML", infilter="writer_pdf_import")
        if format != "docx":
            docx_path = soffice_convert(docx_path, tmp / "out", format)
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(docx_path, output)
    return [output]


def _pdf2docx(src: Path, target: Path, stream_tables: bool) -> None:
    import logging

    from pdf2docx import Converter

    logging.getLogger().setLevel(logging.ERROR)
    cv = Converter(str(src))
    try:
        cv.convert(str(target), multi_processing=False, parse_stream_table=stream_tables)
    finally:
        cv.close()


def _pdf2docx_with_timeout(src: Path, target: Path, seconds: float) -> bool:
    """Full-quality conversion in a child process that is killed if it runs over budget."""
    import multiprocessing as mp

    proc = mp.get_context("fork").Process(target=_pdf2docx, args=(src, target, True))
    proc.start()
    proc.join(seconds)
    if proc.is_alive():
        proc.kill()
        proc.join()
        target.unlink(missing_ok=True)
        return False
    if proc.exitcode != 0 or not target.exists():
        raise ToolError("pdf2docx failed on this file – try --engine libreoffice")
    return True


@tool(
    name="pdf-to-powerpoint", title="PDF to PowerPoint", category="Convert from PDF",
    help="Turn your PDF files into easy to edit PPT and PPTX slideshows.",
    aliases=("pdf-to-pptx", "pdf-to-ppt"), output="pptx", requires=("soffice",),
    opts=[
        Opt("mode", "choice", "image", "'image' = pixel-perfect slides (text kept as speaker notes); "
            "'editable' = LibreOffice import with editable text boxes.", ["image", "editable"]),
        Opt("dpi", "int", 200, "Image mode: slide image resolution."),
    ],
)
def pdf_to_powerpoint(inputs, output: Path, mode="image", dpi=200):
    src = Path(inputs[0])
    if mode == "editable":
        with tempdir() as tmp:
            produced = soffice_convert(src, tmp, "pptx:Impress MS PowerPoint 2007 XML", infilter="impress_pdf_import")
            shutil.move(produced, output)
        return [output]

    from pptx import Presentation
    from pptx.util import Emu

    doc = open_pdf(src)
    prs = Presentation()
    first = doc[0].rect
    prs.slide_width, prs.slide_height = Emu(int(first.width * 12700)), Emu(int(first.height * 12700))
    blank = prs.slide_layouts[6]
    for page in doc:
        slide = prs.slides.add_slide(blank)
        png = page.get_pixmap(dpi=dpi, alpha=False).tobytes("png")
        slide.shapes.add_picture(io.BytesIO(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        text = page.get_text().strip()
        if text:
            slide.notes_slide.notes_text_frame.text = text
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return [output]


@tool(
    name="pdf-to-excel", title="PDF to Excel", category="Convert from PDF",
    help="Pull data straight from PDFs into Excel spreadsheets – tables are detected automatically.",
    aliases=("pdf-to-xlsx",), output="xlsx",
    opts=[
        Opt("layout", "choice", "sheet-per-page", "One sheet per page (tables + text in reading order), every table on "
            "its own sheet (tables only), or everything on one sheet.", ["sheet-per-page", "sheet-per-table", "single-sheet"]),
        Opt("pages", "str", "all", "Pages to convert."),
        Opt("text", "bool", True, "Also export the text around tables (titles, notes) – ignored for sheet-per-table."),
    ],
)
def pdf_to_excel(inputs, output: Path, layout="sheet-per-page", pages="all", text=True):
    from openpyxl import Workbook

    doc = open_pdf(inputs[0])
    wb = Workbook()
    wb.remove(wb.active)
    single = wb.create_sheet("Data") if layout == "single-sheet" else None
    tables_found = 0
    for pno in parse_pages(pages, doc.page_count):
        items = page_layout(doc[pno])
        tables_found += sum(1 for kind, _ in items if kind == "table")
        if layout == "sheet-per-table":
            for ti, rows in enumerate((r for k, r in items if k == "table"), 1):
                _write_items(wb.create_sheet(f"P{pno + 1} T{ti}"[:31]), [("table", rows)])
            continue
        if not text:
            items = [it for it in items if it[0] == "table"]
        if items:
            _write_items(single or wb.create_sheet(f"Page {pno + 1}"), items)
    if not wb.sheetnames:  # sheet-per-table on a document without tables: keep its text
        for pno in parse_pages(pages, doc.page_count):
            items = [it for it in page_layout(doc[pno]) if it[0] == "text"]
            if items:
                _write_items(wb.create_sheet(f"Page {pno + 1}"), items)
    if not wb.sheetnames or all(ws.max_row <= 1 and ws["A1"].value is None for ws in wb):
        raise ToolError("no tables or text found (scanned PDF? run `localpdf ocr` first)")
    print(f"{tables_found} table(s) detected")
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return [output]


def _write_items(ws, items):
    """Append tables (bold header, typed numbers) and text lines (grey italics) to a worksheet."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    widths: dict[int, int] = {}
    prev = None
    for kind, content in items:
        if ws.max_row > 1 and (kind == "table" or prev == "table"):
            ws.append([])  # blank row around tables
        if kind == "text":
            ws.append([content])
            ws.cell(ws.max_row, 1).font = Font(italic=True, color="555555")
        else:
            for r, row in enumerate(content):
                parsed = [_number(c) for c in row]
                ws.append([v for v, _ in parsed])
                for col, (cell, (_, fmt)) in enumerate(zip(ws[ws.max_row], parsed), 1):
                    if fmt:
                        cell.number_format = fmt
                    if r == 0:
                        cell.font = Font(bold=True)
                    widths[col] = max(widths.get(col, 8), min(60, len(str(row[col - 1] or "")) + 2))
        prev = kind
    for col, w in widths.items():  # several pages can share a sheet: only ever widen
        dim = ws.column_dimensions[get_column_letter(col)]
        dim.width = max(w, dim.width or 0)


_NUM = r"\d{1,3}(?:[,\u00a0\u202f ]\d{3})+(?:\.\d+)?|(?:0|[1-9]\d*)(?:\.\d+)?"  # 1,234 / 1 234 / 1234.5
_CURRENCY = "$€£¥₹"


def _number(value):
    """Numeric-looking cell → (number, Excel format); anything else → (text, None).

    Handles thousands separators, percentages, currency symbols and accounting negatives
    "(1,234.00)". IDs with leading zeros ("007") stay text.
    """
    import re

    if value is None:
        return None, None
    s = str(value).strip()
    if not s:
        return None, None
    m = re.fullmatch(rf"(\(?)(-?)\s*([{_CURRENCY}]?)\s*({_NUM})\s*([{_CURRENCY}%]?)(\)?)", s)
    if not m or bool(m.group(1)) != bool(m.group(6)):
        return s, None
    paren, minus, pre, num, post, _ = m.groups()
    v = float(re.sub(r"[,\u00a0\u202f ]", "", num))
    if paren or minus:
        v = -v
    decimals = len(num.split(".")[1]) if "." in num else 0
    base = ("#,##0" if re.search(r"[,\u00a0\u202f ]", num) else "0") + ("." + "0" * decimals if decimals else "")
    if post == "%":
        return round(v / 100, decimals + 2), "0" + ("." + "0" * decimals if decimals else "") + "%"
    if pre or post:
        sym = pre or post
        return v, f'"{sym}"{base}' if pre else f'{base} "{sym}"'
    if decimals == 0 and base == "0":
        return int(v), None
    return (int(v) if decimals == 0 else v), base


@tool(
    name="pdf-to-pdfa", title="PDF to PDF/A", category="Convert from PDF",
    help="Transform your PDF to PDF/A, the ISO-standardized version of PDF for long-term archiving.",
    aliases=("pdfa",), requires=("gs", "ocrmypdf"),
    opts=[Opt("level", "choice", "2b", "PDF/A conformance level.", ["1b", "2b", "3b"])],
)
def pdf_to_pdfa(inputs, output: Path, level="2b"):
    src = Path(inputs[0])
    if which("ocrmypdf"):
        run([which("ocrmypdf"), "--skip-text", "--output-type", f"pdfa-{level[0]}", "--optimize", "0", src, output],
            timeout=1800)
        return [output]
    gs = ghostscript()
    if not gs:
        raise ToolError("PDF/A conversion needs Ghostscript (apt install ghostscript)")
    with tempdir() as tmp:
        icc = _find_icc()
        defs = tmp / "PDFA_def.ps"
        defs.write_text(PDFA_DEF.replace("@ICC@", icc) if icc else PDFA_DEF_NO_ICC)
        cmd = [gs, f"-dPDFA={level[0]}", "-dBATCH", "-dNOPAUSE", "-dQUIET", "-dNOOUTERSAVE",
               "-sColorConversionStrategy=RGB", "-sProcessColorModel=DeviceRGB", "-dPDFACompatibilityPolicy=1",
               "-sDEVICE=pdfwrite", f"-sOutputFile={output}"]
        if icc:
            cmd.append(f"--permit-file-read={icc}")
        run(cmd + [defs, src], timeout=1800)
    return [output]


def _find_icc() -> str | None:
    import glob

    for pattern in ("/usr/share/color/icc/ghostscript/srgb.icc", "/usr/share/ghostscript/*/iccprofiles/srgb.icc",
                    "/usr/share/color/icc/*/sRGB*.icc", "/usr/share/color/icc/sRGB*.icc"):
        hits = glob.glob(pattern)
        if hits:
            return hits[0]
    return None


PDFA_DEF = """%!
/ICCProfile (@ICC@) def
[/Title (Archived document) /DOCINFO pdfmark
[/_objdef {icc_PDFA} /type /stream /OBJ pdfmark
[{icc_PDFA} <</N 3>> /PUT pdfmark
[{icc_PDFA} ICCProfile (r) file /PUT pdfmark
[/_objdef {OutputIntent_PDFA} /type /dict /OBJ pdfmark
[{OutputIntent_PDFA} <</Type /OutputIntent /S /GTS_PDFA1 /DestOutputProfile {icc_PDFA} /OutputConditionIdentifier (sRGB)>> /PUT pdfmark
[{Catalog} <</OutputIntents [ {OutputIntent_PDFA} ]>> /PUT pdfmark
"""
PDFA_DEF_NO_ICC = "%!\n[/Title (Archived document) /DOCINFO pdfmark\n"


# ------------------------------------------------------------------ extras
@tool(
    name="pdf-to-text", title="PDF to Text", category="Extras", extra=True,
    help="Extract all text (reading order preserved), optionally keeping the page layout.",
    aliases=("pdf-to-txt",), output="txt",
    opts=[
        Opt("layout", "bool", False, "Keep the physical layout (columns stay side by side)."),
        Opt("pages", "str", "all", "Pages to extract."),
    ],
)
def pdf_to_text(inputs, output: Path, layout=False, pages="all"):
    doc = open_pdf(inputs[0])
    parts = []
    for pno in parse_pages(pages, doc.page_count):
        page = doc[pno]
        if layout and which("pdftotext"):
            res = run([which("pdftotext"), "-layout", "-f", pno + 1, "-l", pno + 1, inputs[0], "-"])
            parts.append(res.stdout)
        else:
            parts.append("\n".join(reading_order(page)))
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\f".join(parts), encoding="utf-8")
    return [output]


@tool(
    name="pdf-to-html", title="PDF to HTML", category="Extras", extra=True,
    help="Convert a PDF into a single self-contained HTML page (text stays selectable, images embedded).",
    output="html",
)
def pdf_to_html(inputs, output: Path):
    doc = open_pdf(inputs[0])
    pages = "\n".join(f"<section class='page'>{p.get_text('html')}</section>" for p in doc)
    title = doc.metadata.get("title") or Path(inputs[0]).stem
    html = f"""<!doctype html><html><head><meta charset="utf-8"><title>{title}</title>
<style>body{{background:#eee;margin:0;padding:20px}} .page{{margin:0 auto 20px;width:fit-content;box-shadow:0 2px 8px #0003}}
.page>div{{background:#fff;position:relative;overflow:hidden}} .page p{{position:absolute;white-space:pre;margin:0}}
.page img{{position:absolute}}</style></head><body>{pages}</body></html>"""
    output.write_text(html, encoding="utf-8")
    return [output]


@tool(
    name="pdf-to-markdown", title="PDF to Markdown", category="Extras", extra=True,
    help="Convert a PDF into Markdown (headings, lists, tables) – great for notes, wikis and LLMs.",
    aliases=("pdf-to-md",), output="md",
)
def pdf_to_markdown(inputs, output: Path):
    try:
        import pymupdf4llm  # type: ignore

        try:  # pymupdf4llm ≥ 1.x layout engine: no OCR here (that's the `ocr` tool), drop running headers/footers
            md = pymupdf4llm.to_markdown(inputs[0], use_ocr=False, header=False, footer=False, show_progress=False)
        except TypeError:
            md = pymupdf4llm.to_markdown(inputs[0])
    except ImportError:
        md = _simple_markdown(open_pdf(inputs[0]))
    output.write_text(md, encoding="utf-8")
    return [output]


def _simple_markdown(doc) -> str:
    """Heading detection by font size relative to the body text size."""
    from collections import Counter

    sizes = Counter()
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            for line in b.get("lines", []):
                for s in line["spans"]:
                    sizes[round(s["size"])] += len(s["text"])
    body = sizes.most_common(1)[0][0] if sizes else 11
    out = []
    for page in doc:
        for b in page.get_text("dict", sort=True)["blocks"]:
            if b.get("type") != 0:
                continue
            text = " ".join("".join(s["text"] for s in line["spans"]) for line in b["lines"]).strip()
            if not text:
                continue
            size = max(s["size"] for line in b["lines"] for s in line["spans"])
            bold = all("Bold" in s["font"] or s["flags"] & 16 for line in b["lines"] for s in line["spans"])
            if size >= body * 1.6:
                out.append(f"# {text}")
            elif size >= body * 1.3:
                out.append(f"## {text}")
            elif (size >= body * 1.1 or bold) and len(text) < 120:
                out.append(f"### {text}")
            elif text[:1] in "•▪●◦-–*":
                out.append(f"- {text[1:].strip()}")
            else:
                out.append(text)
        out.append("\n---\n")
    return "\n\n".join(out).rstrip("-\n ") + "\n"
