"""Generate the synthetic sample documents in examples/inputs/ (fixed seed → reproducible).

    python examples/make_inputs.py

The real-world PDFs in examples/inputs/realworld/ are public-domain US government documents
(taken from pdfplumber's test corpus) and are not generated.
"""
from __future__ import annotations

import io
import math
import random
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = Path(__file__).parent / "inputs"
rng = random.Random(2026)

FIRST = ["Amélie", "Noah", "Sofia", "Lucas", "Maya", "Ethan", "Chloé", "Omar", "Lena", "Hiro", "Inès", "Jonas"]
LAST = ["Martin", "Okafor", "Rossi", "Nguyen", "Dubois", "Schmidt", "García", "Kowalski", "Tanaka", "Dos Santos"]
PRODUCTS = ["Solar inverter", "Battery pack 10 kWh", "Mounting rail", "Smart meter", "Installation (hours)",
            "Cable kit", "Monitoring licence", "Heat pump", "Charging station", "Maintenance plan"]
WORDS = ("energy storage network project customer quarterly growth market solar install region revenue team "
         "deployment pilot efficiency forecast supplier contract budget pricing risk roadmap capacity").split()


def sentence(n=14):
    w = [rng.choice(WORDS) for _ in range(n)]
    return " ".join(w).capitalize() + "."


def paragraph(k=5):
    return " ".join(sentence(rng.randint(9, 18)) for _ in range(k))


def person():
    return f"{rng.choice(FIRST)} {rng.choice(LAST)}"


def font(size, bold=False):
    name = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    for base in ("/usr/share/fonts/truetype/dejavu/", "/usr/share/fonts/dejavu/", "/usr/share/fonts/TTF/"):
        if Path(base, name).exists():
            return ImageFont.truetype(str(Path(base, name)), size)
    return ImageFont.load_default()


def chart_png(path=None, w=900, h=420, title="Installed capacity (MW)"):
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((30, 18), title, fill="#222", font=font(26, True))
    vals = [rng.randint(20, 100) for _ in range(8)]
    bw = (w - 120) / len(vals)
    for i, v in enumerate(vals):
        x0 = 80 + i * bw + 10
        y0 = h - 50 - v * 3
        d.rectangle([x0, y0, x0 + bw - 20, h - 50], fill=["#1f5fd1", "#e5322d", "#2e9e44", "#f0a30a"][i % 4])
        d.text((x0, h - 42), f"Q{i % 4 + 1}/{24 + i // 4}", fill="#444", font=font(16))
        d.text((x0 + 5, y0 - 22), str(v), fill="#222", font=font(16))
    d.line([70, h - 50, w - 20, h - 50], fill="#333", width=2)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    if path:
        path.write_bytes(buf.getvalue())
    return buf.getvalue()


# ------------------------------------------------------------------ Word
def make_docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Carlito"
    doc.styles["Normal"].font.size = Pt(11)
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "Helios Energy Ltd. — Internal report"
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):  # PAGE field
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        else:
            instr = OxmlElement("w:instrText")
            instr.text = text
            run._r.append(instr)

    doc.add_heading("Quarterly Operations Report", 0)
    p = doc.add_paragraph()
    p.add_run(f"Prepared by {person()} · Reviewed by {person()}").italic = True
    doc.add_heading("1. Executive summary", 1)
    doc.add_paragraph(paragraph(6))
    for _ in range(4):
        doc.add_paragraph(sentence(10), style="List Bullet")
    doc.add_heading("2. Results by region", 1)
    doc.add_paragraph(paragraph(4))
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, h in zip(table.rows[0].cells, ["Region", "Installs", "Revenue (€)", "Growth", "Owner"]):
        cell.text = h
    for region in ["North", "South", "East", "West", "Islands", "Export"]:
        row = table.add_row().cells
        row[0].text = region
        row[1].text = str(rng.randint(40, 900))
        row[2].text = f"{rng.randint(50_000, 2_000_000):,}"
        row[3].text = f"{rng.uniform(-8, 35):.1f}%"
        row[4].text = person()
    doc.add_paragraph()
    doc.add_picture(io.BytesIO(chart_png()), width=Cm(15))
    cap = doc.add_paragraph("Figure 1 – Installed capacity per quarter")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_page_break()
    doc.add_heading("3. Risks and next steps", 1)
    for _ in range(5):
        doc.add_paragraph(sentence(8), style="List Number")
    doc.add_heading("3.1 Detailed analysis", 2)
    for _ in range(5):
        doc.add_paragraph(paragraph(5))
    q = doc.add_paragraph(f"“{sentence(16)}” — {person()}")
    q.paragraph_format.left_indent = Cm(1.5)
    q.runs[0].italic = True
    doc.save(OUT / "report.docx")


# ------------------------------------------------------------------ Excel
def make_xlsx():
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.merge_cells("A1:F1")
    ws["A1"] = "Invoice register – FY2026"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F5FD1")
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = ["Invoice", "Customer", "Product", "Qty", "Unit price", "Total"]
    ws.append(headers)
    thin = Side(style="thin", color="BBBBBB")
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEFB")
    for i in range(1, 41):
        r = i + 2
        ws.append([f"INV-{2026_000 + i}", person(), rng.choice(PRODUCTS), rng.randint(1, 25),
                   round(rng.uniform(15, 4800), 2), f"=D{r}*E{r}"])
        ws[f"E{r}"].number_format = '#,##0.00 "€"'
        ws[f"F{r}"].number_format = '#,##0.00 "€"'
        for c in ws[r]:
            c.border = Border(bottom=thin)
    ws.append(["", "", "", "", "TOTAL", "=SUM(F3:F42)"])
    ws["F43"].font = Font(bold=True)
    ws["F43"].number_format = '#,##0.00 "€"'
    for col, w in zip("ABCDEF", (14, 22, 24, 6, 14, 16)):
        ws.column_dimensions[col].width = w
    ws.print_options.gridLines = False
    ws.page_setup.fitToWidth = 1

    s2 = wb.create_sheet("Summary")
    s2.append(["Month", "Revenue", "Costs", "Margin %"])
    for m, name in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 2):
        s2.append([name, rng.randint(80_000, 240_000), rng.randint(50_000, 160_000), f"=(B{m}-C{m})/B{m}"])
        s2[f"D{m}"].number_format = "0.0%"
    chart = BarChart()
    chart.title = "Revenue vs costs"
    chart.add_data(Reference(s2, min_col=2, max_col=3, min_row=1, max_row=13), titles_from_data=True)
    chart.set_categories(Reference(s2, min_col=1, min_row=2, max_row=13))
    s2.add_chart(chart, "F2")
    wb.save(OUT / "invoices.xlsx")

    # plain CSV too
    lines = ["date,station,temperature_c,humidity_pct,wind_kmh"]
    for d in range(1, 31):
        lines.append(f"2026-09-{d:02d},Lyon-{rng.randint(1, 4)},{rng.uniform(8, 27):.1f},{rng.randint(35, 95)},{rng.uniform(0, 45):.1f}")
    (OUT / "weather.csv").write_text("\n".join(lines) + "\n")


# ------------------------------------------------------------------ PowerPoint
def make_pptx():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Helios Energy – Series B"
    s.placeholders[1].text = f"Investor update · {person()}, CEO"

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Why now?"
    tf = s.placeholders[1].text_frame
    tf.text = sentence(9)
    for _ in range(4):
        p = tf.add_paragraph()
        p.text = sentence(8)
        p.level = rng.choice([0, 1])

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Traction"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3", "Q4"]
    data.add_series("2025", [rng.randint(10, 60) for _ in range(4)])
    data.add_series("2026", [rng.randint(40, 120) for _ in range(4)])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.2), data)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(2.2), Inches(3.8), Inches(2.4))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xE5, 0x32, 0x2D)
    box.text_frame.text = f"+{rng.randint(120, 260)}% YoY"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    box.text_frame.paragraphs[0].runs[0].font.bold = True

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Team"
    rows, cols = 5, 3
    tbl = s.shapes.add_table(rows, cols, Inches(0.8), Inches(1.7), Inches(11.5), Inches(3.5)).table
    for c, h in enumerate(["Name", "Role", "Previously"]):
        tbl.cell(0, c).text = h
    for r in range(1, rows):
        for c, v in enumerate([person(), rng.choice(["CTO", "COO", "VP Sales", "Head of Data"]),
                               rng.choice(["Tesla", "EDF", "Siemens", "Enphase", "Startup founder"])]):
            tbl.cell(r, c).text = v

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(io.BytesIO(chart_png(title="Pipeline by quarter")), Inches(1.5), Inches(0.8), width=Inches(10))
    tb = s.shapes.add_textbox(Inches(1.5), Inches(6.2), Inches(10), Inches(0.8))
    tb.text_frame.text = "Thank you – questions?"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    prs.save(OUT / "pitch.pptx")


# ------------------------------------------------------------------ HTML website
def make_site():
    site = OUT / "website"
    (site / "assets").mkdir(parents=True, exist_ok=True)
    chart_png(site / "assets" / "chart.png", title="Monthly active users (k)")
    (site / "assets" / "logo.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="160" height="40" viewBox="0 0 160 40">'
        '<circle cx="20" cy="20" r="16" fill="#f0a30a"/><text x="44" y="27" font-family="sans-serif" '
        'font-size="20" font-weight="bold" fill="#1d2433">Helios</text></svg>')
    (site / "assets" / "style.css").write_text("""
:root { --brand: #1f5fd1; --ink: #1d2433; --muted: #5b6475; }
* { box-sizing: border-box; }
body { margin: 0; font-family: Inter, "DejaVu Sans", sans-serif; color: var(--ink); line-height: 1.55; }
header { display: flex; align-items: center; justify-content: space-between; padding: 18px 40px; border-bottom: 1px solid #e3e6ee; }
nav a { margin-left: 22px; color: var(--muted); text-decoration: none; }
.hero { padding: 56px 40px; background: linear-gradient(120deg, #1f5fd1, #6a3fd1); color: white; }
.hero h1 { font-size: 42px; margin: 0 0 10px; }
.grid { display: grid; grid-template-columns: repeat(3, 1fr); gap: 20px; padding: 32px 40px; }
.card { border: 1px solid #e3e6ee; border-radius: 12px; padding: 20px; box-shadow: 0 2px 8px rgba(0,0,0,.06); }
.card h3 { margin-top: 0; color: var(--brand); }
table { border-collapse: collapse; width: calc(100% - 80px); margin: 0 40px 30px; }
th, td { border-bottom: 1px solid #e3e6ee; padding: 8px 10px; text-align: left; }
th { background: #f3f6fd; }
td.num { text-align: right; font-variant-numeric: tabular-nums; }
.figure { padding: 0 40px; display: flex; gap: 24px; align-items: center; }
.figure img { width: 55%; border: 1px solid #e3e6ee; border-radius: 8px; }
canvas { border: 1px solid #e3e6ee; border-radius: 8px; }
footer { padding: 20px 40px; color: var(--muted); font-size: 13px; border-top: 1px solid #e3e6ee; }
@media print { .hero { -webkit-print-color-adjust: exact; print-color-adjust: exact; } nav { display: none; } }
""")
    cards = "".join(f'<div class="card"><h3>{t}</h3><p>{sentence(16)}</p></div>'
                    for t in ("Store", "Monitor", "Optimize"))
    rows = "".join(f"<tr><td>{person()}</td><td>{rng.choice(PRODUCTS)}</td><td class='num'>{rng.randint(1, 40)}</td>"
                   f"<td class='num'>{rng.uniform(100, 9000):,.2f} €</td></tr>" for _ in range(8))
    (site / "index.html").write_text(f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Helios Energy – Dashboard</title>
<link rel="stylesheet" href="assets/style.css"></head>
<body>
<header><img src="assets/logo.svg" alt="Helios"><nav><a href="#">Product</a><a href="#">Pricing</a><a href="#">Docs</a></nav></header>
<section class="hero"><h1>Clean energy, managed locally.</h1><p>{sentence(20)}</p></section>
<section class="grid">{cards}</section>
<h2 style="padding:0 40px">Latest orders</h2>
<table><thead><tr><th>Customer</th><th>Product</th><th>Qty</th><th>Amount</th></tr></thead><tbody>{rows}</tbody></table>
<div class="figure"><img src="assets/chart.png" alt="chart"><canvas id="c" width="360" height="240"></canvas></div>
<p style="padding:0 40px" id="generated"></p>
<footer>© 2026 Helios Energy · Rendered from a local file – no web server involved.</footer>
<script>
  // drawn by JavaScript at load time: only a real browser engine will show it
  const c = document.getElementById('c').getContext('2d');
  const vals = [12, 19, 7, 15, 22, 30];
  vals.forEach((v, i) => {{ c.fillStyle = ['#1f5fd1', '#e5322d', '#2e9e44'][i % 3]; c.fillRect(20 + i * 55, 220 - v * 6, 40, v * 6); }});
  c.fillStyle = '#1d2433'; c.font = 'bold 16px sans-serif'; c.fillText('Canvas chart (JS)', 20, 24);
  document.getElementById('generated').textContent = 'This sentence was inserted by JavaScript.';
</script>
</body></html>
""")
    # a standalone single-file page with print CSS, page breaks and @page rules
    (OUT / "invoice.html").write_text(f"""<!doctype html><html><head><meta charset="utf-8"><title>Invoice</title>
<style>
@page {{ size: A4; margin: 18mm; }}
body {{ font-family: "DejaVu Serif", serif; font-size: 11pt; color: #222; }}
h1 {{ letter-spacing: 2px; color: #e5322d; }}
.meta {{ display: flex; justify-content: space-between; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 18px; }}
th {{ background: #222; color: #fff; padding: 6px; text-align: left; }}
td {{ padding: 6px; border-bottom: 1px solid #ddd; }}
.right {{ text-align: right; }} .total {{ font-size: 14pt; font-weight: bold; }}
.terms {{ page-break-before: always; }}
</style></head><body>
<h1>INVOICE</h1>
<div class="meta"><div><b>Helios Energy Ltd.</b><br>12 rue des Lilas<br>69003 Lyon, France</div>
<div class="right">Invoice <b>INV-2026-0417</b><br>Date: 2026-10-14<br>Due: 2026-11-13</div></div>
<p><b>Bill to:</b> {person()}, {rng.randint(1, 99)} avenue Jean Jaurès, Paris</p>
<table><tr><th>Description</th><th class="right">Qty</th><th class="right">Unit</th><th class="right">Amount</th></tr>
{"".join(f"<tr><td>{p}</td><td class='right'>{q}</td><td class='right'>{u:,.2f} €</td><td class='right'>{q * u:,.2f} €</td></tr>"
         for p, q, u in ((rng.choice(PRODUCTS), rng.randint(1, 6), rng.uniform(40, 3000)) for _ in range(7)))}
</table>
<p class="right total">Total due: see attached statement</p>
<div class="terms"><h2>Terms &amp; conditions</h2>{"".join(f"<p>{paragraph(4)}</p>" for _ in range(6))}</div>
</body></html>""")


# ------------------------------------------------------------------ an HTML page designed as a sheet of paper
def make_resume():
    """A CV laid out as a 210×297 mm card – with the classic print traps: a responsive breakpoint wider
    than A4 (so printing switches to the mobile layout) and more content than one page holds."""
    name = person()
    skills = "".join(f"<li>{w}</li>" for w in ("Python", "Linux", "Docker", "SQL", "Git", "Networking", "SIEM", "Bash"))
    jobs = "".join(
        f"<div class='item'><h3>{rng.choice(['Engineer', 'Analyst', 'Developer', 'Project lead'])}</h3>"
        f"<div class='meta'><span>{rng.choice(['Helios Energy', 'Nordlicht AG', 'Atelier Rossi'])}</span>"
        f"<span>{2019 + i} – {2020 + i}</span></div><ul>"
        + "".join(f"<li>{sentence(rng.randint(10, 16))}</li>" for _ in range(3)) + "</ul></div>" for i in range(6))
    (OUT / "resume.html").write_text(f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>{name} – CV</title><style>
body{{margin:0;background:#fff;color:#1d2a38;font:400 16px/1.5 "Source Sans 3","DejaVu Sans",sans-serif}}
a{{color:#2a62a0;text-decoration:none}}
.page{{width:210mm;margin:0 auto}}
.sheet{{background:#fff;border:1px solid #e3e9f0;width:210mm;min-height:297mm;margin:0 auto}}
.top{{padding:10mm 10mm 7mm;border-bottom:3px solid #2a62a0}} .top h1{{margin:0;font-size:2rem}}
.cols{{display:grid;grid-template-columns:60mm 1fr}}
aside{{background:#f2f6fb;margin:8mm 0 8mm 8mm;padding:14px;border-radius:8px;font-size:.8rem}}
aside ul{{list-style:none;padding:0;display:flex;flex-wrap:wrap;gap:.3rem}}
aside li{{background:#fff;border:1px solid #e3e9f0;border-radius:6px;padding:.1rem .45rem}}
main{{padding:8mm 10mm;font-size:.85rem}} h2{{color:#2a62a0;font-size:1.05rem}}
.item{{border-left:2px solid #e3e9f0;padding:0 0 .8rem 1.1rem}} .item h3{{margin:0;font-size:.92rem}}
.meta{{display:flex;justify-content:space-between;color:#66768a;font-size:.78rem}}
@media screen{{body{{background:#e7ebf1}} .page{{padding:8mm 0}} .sheet{{border-radius:6px;box-shadow:0 6px 28px rgba(29,42,56,.12)}}}}
@media (max-width:230mm){{.page{{width:100%}} .sheet{{width:100%}} .cols{{grid-template-columns:1fr}}}}
</style></head><body><div class="page"><article class="sheet">
<header class="top"><h1>{name}</h1><p>Security engineer · open to new opportunities</p></header>
<div class="cols"><aside><h2>Contact</h2><p><a href="mailto:jane@example.org">jane@example.org</a><br>
<a href="https://example.org/portfolio">example.org/portfolio</a></p><h2>Skills</h2><ul>{skills}</ul>
<h2>Languages</h2><p>English – native<br>French – C1<br>Korean – A1</p></aside>
<main><h2>Profile</h2><p>{paragraph(4)}</p><h2>Experience</h2>{jobs}</main></div>
</article></div></body></html>""")


# ------------------------------------------------------------------ Markdown, text, RTF
def make_text_docs():
    rows = "\n".join(f"| {rng.choice(PRODUCTS)} | {rng.randint(1, 99)} | {rng.uniform(1, 999):.2f} |" for _ in range(5))
    (OUT / "README-sample.md").write_text(f"""# Project Helios – setup notes

> {sentence(14)}

## Requirements

- Python 3.10+
- [x] Ghostscript installed
- [ ] LibreOffice configured

## Inventory

| Item | Qty | Price |
|------|----:|------:|
{rows}

## Install

```bash
pip install helios-agent
helios init --site lyon-3 --capacity 250kW
```

### Notes

{paragraph(5)}

1. {sentence(8)}
2. {sentence(8)}
3. {sentence(8)}

Inline `code`, **bold**, *italic* and a [link](https://example.org).

![chart](website/assets/chart.png)
""")
    levels = ["INFO", "INFO", "INFO", "WARN", "ERROR", "DEBUG"]
    lines = [f"2026-10-29T{h:02d}:{m:02d}:{s:02d}Z {rng.choice(levels):5} [worker-{rng.randint(1, 8)}] "
             f"{rng.choice(['request', 'sync', 'battery', 'meter', 'auth'])}: {sentence(rng.randint(5, 20))}"
             for h, m, s in sorted((rng.randint(0, 23), rng.randint(0, 59), rng.randint(0, 59)) for _ in range(260))]
    (OUT / "server.log").write_text("\n".join(lines) + "\n")
    (OUT / "letter.rtf").write_text(
        r"{\rtf1\ansi\deff0{\fonttbl{\f0 Times New Roman;}{\f1 Arial;}}{\colortbl;\red229\green50\blue45;}"
        r"\f1\fs32\b\cf1 Formal letter\b0\cf0\fs24\par\par\f0 Dear " + person() + r",\par\par "
        + paragraph(5) + r"\par\par\i " + sentence(10) + r"\i0\par\par Kind regards,\par " + person() + r"\par}")


# ------------------------------------------------------------------ images
def render_paper(lines=24, w=1240, h=1754):
    """A page of text as a 150-dpi-ish image (used for photos and scans)."""
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((110, 110), "DELIVERY NOTE  No. 58213", fill="black", font=font(46, True))
    d.text((110, 190), f"Customer: {person()}", fill="black", font=font(30))
    y = 280
    truth = ["DELIVERY NOTE No. 58213"]
    for _ in range(lines):
        line = sentence(rng.randint(6, 9))[:62]
        d.text((110, y), line, fill="black", font=font(28))
        truth.append(line)
        y += 52
    return img, truth


def make_images():
    # phone photo of a document: rotated, uneven lighting, slightly blurry, on a desk
    paper, _ = render_paper()
    paper = paper.rotate(3.5, expand=True, fillcolor=(120, 90, 60), resample=Image.BICUBIC)
    desk = Image.new("RGB", (paper.width + 160, paper.height + 160), (120, 90, 60))
    desk.paste(paper, (80, 80))
    shade = Image.new("L", desk.size)
    sd = ImageDraw.Draw(shade)
    for i in range(desk.width):  # light falls off from left to right
        sd.line([(i, 0), (i, desk.height)], fill=int(255 - 110 * i / desk.width))
    photo = Image.composite(desk, Image.new("RGB", desk.size, "black"), shade).filter(ImageFilter.GaussianBlur(1.2))
    photo.save(OUT / "phone-photo-of-document.jpg", quality=82)

    # colourful photo-like image
    w, h = 1600, 1000
    img = Image.new("RGB", (w, h))
    px = img.load()
    for y in range(h):
        for x in range(0, w):
            px[x, y] = (int(40 + 180 * y / h), int(120 + 80 * math.sin(x / 140)), int(220 - 150 * y / h))
    d = ImageDraw.Draw(img)
    d.ellipse([1150, 120, 1400, 370], fill=(255, 220, 80))
    d.polygon([(0, 1000), (500, 450), (900, 1000)], fill=(40, 70, 50))
    d.polygon([(600, 1000), (1100, 520), (1600, 1000)], fill=(30, 55, 40))
    img.save(OUT / "landscape.webp", quality=90)
    img.resize((800, 500)).save(OUT / "landscape-small.jpg", quality=90)

    # transparent PNG logo
    logo = Image.new("RGBA", (600, 300), (0, 0, 0, 0))
    d = ImageDraw.Draw(logo)
    d.ellipse([20, 20, 280, 280], fill=(240, 163, 10, 255))
    d.text((300, 110), "Helios", fill=(29, 36, 51, 255), font=font(80, True))
    logo.save(OUT / "logo-transparent.png")

    # portrait photo with EXIF orientation = 6 (camera held sideways)
    side = Image.new("RGB", (1200, 800), (200, 225, 245))
    ImageDraw.Draw(side).text((60, 340), "This text should end up UPRIGHT", fill="black", font=font(56, True))
    exif = Image.Exif()
    exif[0x0112] = 6  # pixels stored turned 90° counter-clockwise; viewers turn them back clockwise
    side.rotate(90, expand=True).save(OUT / "exif-rotated.jpg", exif=exif, quality=90)

    # 3-page TIFF fax
    pages = []
    for i in range(3):
        pg, _ = render_paper(lines=10)
        pages.append(pg.convert("1"))
        ImageDraw.Draw(pages[-1]).text((110, 1600), f"FAX page {i + 1}/3", fill=0, font=font(30))
    pages[0].save(OUT / "fax-3-pages.tiff", save_all=True, append_images=pages[1:], compression="group4")


# ------------------------------------------------------------------ scanned PDF (for OCR) + ground truth
def make_scan():
    import pymupdf as fitz

    doc = fitz.open()
    truth_all = []
    for _ in range(2):
        img, truth = render_paper(lines=26)
        truth_all += truth
        img = img.rotate(-0.8, fillcolor="white", resample=Image.BICUBIC).convert("L")
        noise = Image.effect_noise(img.size, 18)
        img = Image.blend(img, noise.convert("L"), 0.08)
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=70)
        page = doc.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(OUT / "scanned-delivery-notes.pdf")
    (OUT / "scanned-delivery-notes.truth.txt").write_text("\n".join(truth_all) + "\n")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for fn in (make_docx, make_xlsx, make_pptx, make_site, make_text_docs, make_images, make_scan, make_resume):
        fn()
        print("✔", fn.__name__)
